#!/usr/bin/env python3
"""Python 3.10+；从标准输入接收 UTF-8 JSON，用 openpyxl 生成配置表。

调用：python 表格生成工具.py --项目目录 项目根目录 [--覆盖]
JSON 从标准输入传入，不需要额外数据文件；--格式 可输出完整输入示例。
保存后自动调用项目的表格检查工具，业务规则以其结果为准。
退出码：0 检查通过，1 检查未通过，2 依赖、输入、写入或工具运行失败。
"""

import argparse
from io import BytesIO
import json
import math
from pathlib import Path
import re
import subprocess
import sys


输入示例 = {
    "文件名": "生成示例.xlsx",
    "分类": "数据配置",
    "工作表": [
        {
            "名称": "|英雄",
            "行": [
                ["ID", "名字", "生命", "启用", "武器", "属性", "备注"],
                ["key", "string", "int", "bool", "string[]", "{int 最小;int 最大}", "string"],
                [1, "示例英雄", None, False, "长剑,盾牌", "1;99", "=这是文本"],
                [2, "第二英雄", 0, True, None, None, "00123"],
            ],
        },
        {
            "名称": "|全局参数",
            "行": [
                ["name", "value", "type", "sign", "description"],
                ["初始金币", None, "int", None, "留空时由导表工具处理"],
            ],
        },
    ],
}


def 输出(内容):
    print(json.dumps(内容, ensure_ascii=False))


def 加载依赖():
    try:
        import openpyxl
    except ModuleNotFoundError as 异常:
        if 异常.name != "openpyxl":
            raise ValueError(f"openpyxl 依赖不完整：缺少 {异常.name}，请先修复安装。") from 异常
        raise ValueError(
            '缺少 openpyxl，已停止生成。请先安装：'
            f'"{sys.executable}" -m pip install openpyxl'
        ) from 异常
    return openpyxl


def 检查键(对象, 必需, 可选, 位置):
    if not isinstance(对象, dict):
        raise ValueError(f"{位置}必须为 JSON 对象")
    缺少 = 必需 - 对象.keys()
    多余 = 对象.keys() - 必需 - 可选
    if 缺少 or 多余:
        raise ValueError(f"{位置}字段不符：缺少 {sorted(缺少)}，不支持 {sorted(多余)}")


def 读取输入():
    # utf-8-sig 同时接受带 BOM 和不带 BOM 的 UTF-8 管道输入。
    输入 = json.loads(sys.stdin.buffer.read().decode("utf-8-sig"))
    检查键(输入, {"文件名", "工作表"}, {"分类"}, "输入")
    文件名 = 输入["文件名"]
    if (not isinstance(文件名, str) or not 文件名.lower().endswith(".xlsx")
            or len(文件名) <= 5 or re.search(r'[\\/:*?"<>|\x00-\x1f]', 文件名)
            or 文件名.startswith(("~$", ".~"))):
        raise ValueError("文件名必须是有效的单个 .xlsx 文件名，不能包含路径或 Office 临时文件前缀")
    if 输入.get("分类", "数据配置") not in ("数据配置", "文案配置"):
        raise ValueError("分类仅支持 数据配置 或 文案配置")
    if not isinstance(输入["工作表"], list) or not 输入["工作表"]:
        raise ValueError("工作表必须为非空列表")
    return 输入


def 创建工作簿(openpyxl, 输入):
    工作簿 = openpyxl.Workbook()
    工作簿.remove(工作簿.active)
    表名集合 = set()
    for 序号, 表 in enumerate(输入["工作表"], 1):
        检查键(表, {"名称", "行"}, set(), f"第 {序号} 张工作表")
        名称, 行列表 = 表["名称"], 表["行"]
        if (not isinstance(名称, str) or not 名称.strip() or len(名称) > 31
                or re.search(r"[\\/*?:\[\]\x00-\x1f]", 名称)):
            raise ValueError(f"第 {序号} 张工作表名称无效，需为 1 至 31 字且不含 Excel 禁用字符")
        if 名称.casefold() in 表名集合:
            raise ValueError(f"工作表名称重复：{名称}")
        表名集合.add(名称.casefold())
        if not isinstance(行列表, list) or not 行列表 or len(行列表) > 1048576:
            raise ValueError(f"{名称} 的行必须为非空列表，最多 1048576 行")
        工作表 = 工作簿.create_sheet(名称)
        列宽 = {}
        for 行号, 行 in enumerate(行列表, 1):
            if not isinstance(行, list) or len(行) > 16384:
                raise ValueError(f"{名称} 第 {行号} 行必须为列表，最多 16384 列")
            for 列号, 值 in enumerate(行, 1):
                if 值 is None or 值 == "":
                    continue
                地址 = f"{名称}!{openpyxl.utils.get_column_letter(列号)}{行号}"
                if not isinstance(值, (str, int, float, bool)):
                    raise ValueError(f"{地址} 只接受字符串、数字、布尔或 null；容器字段需整理成文本")
                if isinstance(值, str) and len(值) > 32767:
                    raise ValueError(f"{地址} 文本超过 Excel 的 32767 字符限制")
                if isinstance(值, float) and not math.isfinite(值):
                    raise ValueError(f"{地址} 不接受 NaN 或 Infinity")
                if isinstance(值, int) and not isinstance(值, bool) and abs(值) >= 10 ** 15:
                    raise ValueError(f"{地址} 超过 15 位的整数请用字符串传入，以免丢失精度")
                try:
                    单元格 = 工作表.cell(行号, 列号, 值)
                except openpyxl.utils.exceptions.IllegalCharacterError as 异常:
                    raise ValueError(f"{地址} 含 Excel 不允许的控制字符") from 异常
                if isinstance(值, str):
                    # 同时防止 = 开头文本和 #N/A 等文本被推断成公式或错误值。
                    单元格.data_type = "s"
                单元格.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
                显示宽度 = max((sum(2 if ord(字) > 127 else 1 for 字 in 一行)
                               for 一行 in str(值).splitlines()), default=0)
                列宽[列号] = max(列宽.get(列号, 12), min(显示宽度 + 2, 48))
        首行 = 行列表[0]
        是配置表 = all(名 in 首行 for 名 in ("name", "value", "type"))
        表头行数 = 1 if 是配置表 or not 名称.startswith("|") else 2
        for 单元格行 in 工作表.iter_rows(min_row=1, max_row=表头行数):
            for 单元格 in 单元格行:
                单元格.font = openpyxl.styles.Font(bold=True)
        工作表.freeze_panes = f"A{表头行数 + 1}"
        for 列号, 宽度 in 列宽.items():
            工作表.column_dimensions[openpyxl.utils.get_column_letter(列号)].width = 宽度
    return 工作簿


def 运行检查(检查工具, 表格目录):
    进程 = subprocess.run(
        [sys.executable, "-B", "-X", "utf8", str(检查工具), "--目录", str(表格目录), "--json"],
        capture_output=True, text=True, encoding="utf-8", check=False,
    )
    try:
        结果 = json.loads(进程.stdout)
    except json.JSONDecodeError as 异常:
        raise ValueError(f"表格检查工具未返回有效 JSON：{进程.stderr.strip() or 进程.stdout.strip()}") from 异常
    if not isinstance(结果, dict) or not isinstance(结果.get("通过"), bool):
        raise ValueError("表格检查工具缺少布尔类型的‘通过’结果")
    if 进程.returncode not in (0, 1):
        raise ValueError(f"表格检查工具运行异常，退出码 {进程.returncode}：{进程.stderr.strip()}")
    return 结果, 0 if 进程.returncode == 0 and 结果["通过"] else 1


def 主函数():
    参数器 = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    参数器.add_argument("--项目目录", type=Path, help="含 4数据配置表格 的项目根目录，非 Godot 工程目录")
    参数器.add_argument("--覆盖", action="store_true", help="显式允许替换整个已有工作簿，输入需包含全部要保留的内容")
    参数器.add_argument("--依赖检查", action="store_true", help="仅检查当前解释器能否加载 openpyxl，不生成文件")
    参数器.add_argument("--格式", action="store_true", help="输出 JSON 输入示例，不生成文件；分类另支持 文案配置")
    参数 = 参数器.parse_args()
    if 参数.格式:
        print(json.dumps(输入示例, ensure_ascii=False, indent=2))
        return 0
    已生成 = None
    try:
        openpyxl = 加载依赖()
        if 参数.依赖检查:
            输出({"依赖可用": True, "Python": sys.executable, "openpyxl": openpyxl.__version__})
            return 0
        if 参数.项目目录 is None:
            raise ValueError("生成时必须传入 --项目目录，并通过标准输入提供 UTF-8 JSON")
        项目目录 = 参数.项目目录.resolve()
        表格目录 = 项目目录 / "4数据配置表格"
        检查工具 = 表格目录 / "表格检查工具.py"
        if not 检查工具.is_file():
            raise ValueError(f"缺少表格检查工具，已停止生成：{检查工具}")
        输入 = 读取输入()
        目标目录 = 表格目录 / "文案表格" if 输入.get("分类") == "文案配置" else 表格目录
        目标路径 = (目标目录 / 输入["文件名"]).resolve()
        if not 目标路径.is_relative_to(表格目录.resolve()):
            raise ValueError("目标路径超出 4数据配置表格 目录")
        if 目标路径.exists() and not 参数.覆盖:
            raise ValueError(f"文件已存在，未写入；授权替换整个工作簿后使用 --覆盖：{目标路径}")
        工作簿 = 创建工作簿(openpyxl, 输入)
        # 先在内存完成序列化，输入或编码错误不会截断已有工作簿。
        with BytesIO() as 缓冲:
            工作簿.save(缓冲)
            目标目录.mkdir(parents=True, exist_ok=True)
            with 目标路径.open("wb" if 参数.覆盖 else "xb") as 文件:
                文件.write(缓冲.getbuffer())
        工作簿.close()
        已生成 = str(目标路径)
        结果, 退出码 = 运行检查(检查工具, 表格目录)
        输出({"已生成": 已生成, "检查": 结果})
        return 退出码
    except (OSError, ValueError, TypeError, OverflowError, ImportError) as 异常:
        输出({"已生成": 已生成, "错误": str(异常) or type(异常).__name__})
        return 2


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    sys.exit(主函数())
